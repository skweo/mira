#!/usr/bin/env python3
"""Plan and audit visual opportunities for Mira contest papers.

This gate asks whether Mira noticed the right figure grammar for available
evidence. It complements figure_portfolio_gate.py: this script finds missed
opportunities such as radar charts, heatmaps, 3D surfaces, hexbin joint
distributions, distribution plots, PPT reasoning diagrams, and Excel-style
summary visuals before final writing.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any


TABLE_SUFFIXES = {".csv", ".tsv", ".xlsx", ".xlsm"}
FIGURE_SUFFIXES = {".pdf", ".png", ".jpg", ".jpeg", ".svg", ".webp"}
PAPER_CANDIDATES = [
    "paper/main_final_checked.tex",
    "paper/main_final_repaired.tex",
    "paper/main_final.tex",
    "paper/main_contest.tex",
    "paper/main.tex",
    "paper/main.typ",
    "paper/main.md",
]

QUESTION_RE = re.compile(r"(?:^|[/_\\-])q([1-9])(?:[/_\\_.-]|$)|问题([一二三四五六七八九])", re.I)
CN_NUMBERS = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}


@dataclass
class TableProfile:
    path: str
    question: str
    rows: int
    cols: int
    headers: list[str]
    numeric_cols: list[str]
    text_cols: list[str]
    sample_values: dict[str, list[str]]
    notes: list[str]


@dataclass
class Opportunity:
    opportunity_id: str
    question: str
    source: str
    priority: str
    role: str
    recommended_visual: str
    recommended_toolchain: str
    reason: str
    data_contract: str
    command_hint: str
    status: str
    realization_evidence: str


@dataclass
class ThreeDCandidate:
    candidate_id: str
    source: str
    question: str
    kind: str
    priority: str
    reason: str
    data_columns: list[str]
    required_companion: str
    status: str = "unresolved"
    artifact: str = ""
    waiver_reason: str = ""
    waiver_evidence: list[str] | None = None
    disposition_issue: str = "no storyboard disposition recorded"


@dataclass
class Finding:
    level: str
    axis: str
    return_phase: str
    finding: str
    recommendation: str
    evidence: str = ""


def main() -> int:
    args = parse_args()
    root = Path(args.root).resolve()
    table_profiles = collect_table_profiles(root)
    visual_text = collect_visual_text(root)
    opportunities = build_opportunities(root, table_profiles, visual_text)
    storyboard = load_json(root / "planning" / "figure_storyboard.json")
    three_d_candidates = build_three_d_candidates(root, table_profiles)
    apply_three_d_dispositions(root, three_d_candidates, storyboard.get("three_d_candidates", []))
    findings, metrics = audit(
        root, table_profiles, opportunities, three_d_candidates, visual_text, args.output_level
    )
    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "root": str(root),
        "verdict": verdict(findings),
        "metrics": metrics,
        "table_profiles": [asdict(item) for item in table_profiles],
        "opportunities": [asdict(item) for item in opportunities],
        "three_d_candidates": [asdict(item) for item in three_d_candidates],
        "findings": [asdict(item) for item in findings],
    }

    if args.write_report:
        out = resolve_path(root, args.write_report)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(markdown(payload), encoding="utf-8")
    if args.write_json:
        out = resolve_path(root, args.write_json)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    emit(payload)
    return 1 if payload["verdict"] == "FAIL" else 0


def collect_table_profiles(root: Path) -> list[TableProfile]:
    table_dir = root / "results" / "tables"
    if not table_dir.exists():
        return []
    profiles: list[TableProfile] = []
    for path in sorted(table_dir.rglob("*")):
        if path.suffix.lower() not in TABLE_SUFFIXES:
            continue
        profile = profile_table(root, path)
        if profile:
            profiles.append(profile)
    return profiles


def profile_table(root: Path, path: Path) -> TableProfile | None:
    rows = read_rows(path)
    if not rows:
        return None
    headers = [str(item).strip() for item in rows[0]]
    if not headers:
        return None
    data = rows[1:]
    numeric_cols: list[str] = []
    text_cols: list[str] = []
    sample_values: dict[str, list[str]] = {}
    for idx, header in enumerate(headers):
        values = [row[idx] if idx < len(row) else "" for row in data]
        nonempty = [value for value in values if str(value).strip()]
        sample_values[header] = [str(value) for value in nonempty[:5]]
        numeric_ratio = sum(1 for value in nonempty if is_number(value)) / max(1, len(nonempty))
        if nonempty and numeric_ratio >= 0.75:
            numeric_cols.append(header)
        else:
            text_cols.append(header)
    notes = []
    if len(data) >= 100:
        notes.append("long_sample_table")
    if any(re.search(r"angle|方向|角度|theta|deg", header, re.I) for header in headers):
        notes.append("angle_or_direction")
    if any(re.search(r"noise|sd|sigma|duration|参数|灵敏|sensitivity|alpha|lambda|cost|权重", header, re.I) for header in headers):
        notes.append("parameter_or_sensitivity")
    if any(re.search(r"probability|prob|概率|rate|success|成功", header, re.I) for header in headers):
        notes.append("probability_response")
    return TableProfile(
        path=rel(root, path),
        question=detect_question(path.name),
        rows=len(data),
        cols=len(headers),
        headers=headers,
        numeric_cols=numeric_cols,
        text_cols=text_cols,
        sample_values=sample_values,
        notes=notes,
    )


def read_rows(path: Path) -> list[list[Any]]:
    if path.suffix.lower() in {".csv", ".tsv"}:
        text = read_text(path)
        if not text.strip():
            return []
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        try:
            sample = text[:4096]
            dialect = csv.Sniffer().sniff(sample)
            delimiter = dialect.delimiter
        except csv.Error:
            pass
        return list(csv.reader(text.splitlines(), delimiter=delimiter))
    return read_xlsx_rows(path)


def read_xlsx_rows(path: Path) -> list[list[Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        return []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        return [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
    except Exception:
        return []


def build_opportunities(root: Path, profiles: list[TableProfile], visual_text: str) -> list[Opportunity]:
    opportunities: list[Opportunity] = []
    for profile in profiles:
        opportunities.extend(opportunities_for_table(root, profile, visual_text))
    opportunities.extend(contextual_opportunities(root, visual_text))
    return dedupe_opportunities(opportunities)


def build_three_d_candidates(root: Path, profiles: list[TableProfile]) -> list[ThreeDCandidate]:
    candidates: list[ThreeDCandidate] = []
    for profile in profiles:
        rows = read_rows(root / profile.path)
        if len(rows) < 2:
            continue
        candidate = three_d_candidate_for_table(profile, rows[1:])
        if candidate:
            candidates.append(candidate)
    return candidates


def three_d_candidate_for_table(profile: TableProfile, rows: list[list[Any]]) -> ThreeDCandidate | None:
    columns = {canonical_header(header): header for header in profile.numeric_cols}
    context = (" ".join(profile.headers) + " " + profile.path).lower()
    x = find_column(columns, "x", "xcoord", "xcoordinate", "coordx", "east", "easting", "longitude", "lon")
    y = find_column(columns, "y", "ycoord", "ycoordinate", "coordy", "north", "northing", "latitude", "lat")
    z = find_column(columns, "z", "zcoord", "zcoordinate", "coordz", "height", "altitude", "elevation")
    time_col = find_column(columns, "t", "time", "timestamp", "step", "timestep", "sequence", "order", "iteration")
    value_col = find_matching_column(
        profile.numeric_cols,
        r"value|scalar|field|temperature|pressure|density|intensity|potential|concentration|response|objective|score|prob|rate|error|cost|值|温度|压力|密度|浓度|响应|目标",
        exclude={x, y, z},
    )
    u = find_column(columns, "u", "vx", "velocityx", "vectorx", "fieldx", "componentx")
    v = find_column(columns, "v", "vy", "velocityy", "vectory", "fieldy", "componenty")
    w = find_column(columns, "w", "vz", "velocityz", "vectorz", "fieldz", "componentz")

    face_cols = [
        col for col in profile.headers
        if re.search(r"face|triangle|vertex[_ -]?[123]|node[_ -]?[123]|面片|三角|顶点", col, re.I)
    ]
    mesh_signal = bool(face_cols) or bool(
        re.search(r"mesh|triang|surface_geometry|faces|vertices|拓扑|面片|三角剖分", context, re.I)
    )
    if mesh_signal and (len(face_cols) >= 3 or all((x, y, z))):
        return make_three_d_candidate(
            profile,
            "mesh_geometry",
            [col for col in [x, y, z, *face_cols] if col],
            "vertex/face or mesh topology fields describe three-dimensional geometry",
        )

    if all((x, y, z, u, v, w)):
        return make_three_d_candidate(
            profile,
            "vector_field_3d",
            [x, y, z, u, v, w],
            "three spatial coordinates and three vector components define a 3D vector field",
        )

    if all((x, y, z, value_col)):
        return make_three_d_candidate(
            profile,
            "scalar_field_3d",
            [x, y, z, value_col],
            "three spatial coordinates plus a scalar value define a 3D field sample",
        )

    if all((x, y, z, time_col)) and profile.rows >= 4:
        return make_three_d_candidate(
            profile,
            "trajectory_3d",
            [time_col, x, y, z],
            "ordered observations contain time and three spatial coordinates",
        )

    if has_two_parameter_response(profile):
        params = parameter_columns(profile)
        response = (response_columns(profile) or [profile.numeric_cols[-1]])[0]
        if len(params) >= 2:
            grid_note = "complete parameter grid" if is_complete_grid(profile, rows, params[0], params[1]) else "scattered or incomplete parameter grid"
            return make_three_d_candidate(
                profile,
                "two_parameter_sensitivity",
                [params[0], params[1], response],
                f"two independent parameters and one response form a {grid_note}",
            )

    if all((x, y, z)):
        if is_complete_grid(profile, rows, x, y) or re.search(r"surface|response.?surface|grid|height.?map|曲面|响应面|网格", context, re.I):
            return make_three_d_candidate(
                profile,
                "surface_grid",
                [x, y, z],
                "x-y samples with a z response form a surface grid or response surface",
            )
        if profile.rows >= 30 or re.search(r"point.?cloud|lidar|scan|sample.?points|点云|激光雷达", context, re.I):
            return make_three_d_candidate(
                profile,
                "point_cloud_3d",
                [x, y, z],
                "many observations provide three independent spatial coordinates",
            )
        return make_three_d_candidate(
            profile,
            "spatial_coordinates",
            [x, y, z],
            "the table contains explicit three-dimensional spatial coordinates",
        )
    return None


def make_three_d_candidate(
    profile: TableProfile, kind: str, data_columns: list[str], reason: str
) -> ThreeDCandidate:
    return ThreeDCandidate(
        candidate_id=slug(f"{Path(profile.path).stem}-3d-{kind}"),
        source=profile.path,
        question=profile.question,
        kind=kind,
        priority="high",
        reason=reason,
        data_columns=data_columns,
        required_companion="contour_projection_slice_or_table",
        waiver_evidence=[],
    )


def apply_three_d_dispositions(
    root: Path, candidates: list[ThreeDCandidate], dispositions: Any
) -> None:
    records = {
        str(item.get("candidate_id", "")).strip(): item
        for item in dispositions if isinstance(item, dict)
    } if isinstance(dispositions, list) else {}
    for candidate in candidates:
        record = records.get(candidate.candidate_id)
        if not record:
            continue
        requested = str(record.get("status", "")).strip().lower()
        artifact = str(record.get("artifact", "")).strip()
        waiver_reason = str(record.get("waiver_reason", "")).strip()
        waiver_evidence = clean_list(record.get("waiver_evidence"))
        candidate.artifact = artifact
        candidate.waiver_reason = waiver_reason
        candidate.waiver_evidence = waiver_evidence
        if requested == "generated":
            artifact_path = resolve_path(root, artifact)
            if (
                artifact
                and not is_placeholder(artifact)
                and artifact_path.suffix.lower() in FIGURE_SUFFIXES
                and artifact_path.is_file()
            ):
                candidate.status = "generated"
                candidate.disposition_issue = ""
            else:
                candidate.disposition_issue = "generated status requires an existing figure artifact"
        elif requested == "waived":
            issue = waiver_issue(waiver_reason, waiver_evidence)
            if not issue:
                candidate.status = "waived"
                candidate.disposition_issue = ""
            else:
                candidate.disposition_issue = issue
        else:
            candidate.disposition_issue = "status must be generated or waived"


def waiver_issue(reason: str, evidence: list[str]) -> str:
    generic = {
        "not needed", "unnecessary", "no need", "not useful", "skip", "waived",
        "不需要", "没必要", "无需", "不用画", "跳过", "不采用",
    }
    if is_placeholder(reason) or len(reason) < 12 or reason.lower() in generic:
        return "waiver requires a specific readability, dimensionality, or data-sufficiency reason"
    if not evidence or any(is_placeholder(item) for item in evidence):
        return "waiver requires concrete comparison, table, projection, or data evidence"
    return ""


def canonical_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def find_column(columns: dict[str, str], *aliases: str) -> str:
    for alias in aliases:
        if alias in columns:
            return columns[alias]
    return ""


def find_matching_column(columns: list[str], pattern: str, exclude: set[str]) -> str:
    return next((column for column in columns if column not in exclude and re.search(pattern, column, re.I)), "")


def is_complete_grid(profile: TableProfile, rows: list[list[Any]], x_col: str, y_col: str) -> bool:
    try:
        x_idx = profile.headers.index(x_col)
        y_idx = profile.headers.index(y_col)
    except ValueError:
        return False
    pairs = {
        (str(row[x_idx]).strip(), str(row[y_idx]).strip())
        for row in rows
        if x_idx < len(row) and y_idx < len(row) and is_number(row[x_idx]) and is_number(row[y_idx])
    }
    xs = {pair[0] for pair in pairs}
    ys = {pair[1] for pair in pairs}
    return len(xs) >= 3 and len(ys) >= 3 and len(pairs) == len(xs) * len(ys)


def opportunities_for_table(root: Path, profile: TableProfile, visual_text: str) -> list[Opportunity]:
    ops: list[Opportunity] = []
    lower_path = profile.path.lower()
    num = profile.numeric_cols
    text = profile.text_cols
    rows = profile.rows
    q = profile.question

    if is_metric_value_table(profile):
        metric_count = rows
        priority = "medium" if 4 <= metric_count <= 10 else "optional"
        ops.append(
            make_opportunity(
                profile,
                "radar",
                q,
                priority,
                "compare",
                "雷达图或归一化指标条形图",
                "Python/matplotlib radar; Excel-compatible normalized metric table",
                "metric-value table can become a compact multi-indicator profile after direction and scale normalization",
                "one metric column and one numeric value column; normalize and state whether larger is better",
                radar_command(profile),
                visual_text,
            )
        )

    if has_scheme_metric_matrix(profile):
        metric_count = len(num)
        row_count = rows
        if row_count <= 6 and metric_count <= 8:
            visual = "雷达图"
            toolchain = "Python/matplotlib radar or Excel radar chart from normalized XLSX/CSV"
            command = radar_command(profile)
        else:
            visual = "平行坐标图"
            toolchain = "Python/matplotlib parallel coordinates; Excel-compatible normalized table"
            command = parallel_command(profile)
        ops.append(
            make_opportunity(
                profile,
                "multi_metric_compare",
                q,
                "high",
                "compare",
                visual,
                toolchain,
                f"{row_count} schemes/items with {metric_count} numeric indicators should not be shown only as a plain table",
                "one item/scheme column and 3-8 numeric metric columns; align metric directions before plotting",
                command,
                visual_text,
            )
        )

    if has_two_parameter_response(profile):
        response_cols = response_columns(profile)
        response = response_cols[0] if response_cols else (num[-1] if num else "")
        ops.append(
            make_opportunity(
                profile,
                "heatmap_surface",
                q,
                "high",
                "validate",
                "二维热力图，必要时配 MATLAB 风格三维曲面/等高线图",
                "Python/matplotlib heatmap/surface; MATLAB/Octave-style response surface when available",
                "two parameter columns and response metrics form a sensitivity surface; a line chart may hide interaction effects",
                "numeric x, y, z columns with a grid or scattered response; record interpolation if surface is not a complete grid",
                surface_command(profile, response),
                visual_text,
            )
        )

    if has_bivariate_sample_cloud(profile):
        x_col, y_col = bivariate_sample_columns(profile)
        ops.append(
            make_opportunity(
                profile,
                "hexbin_joint",
                q,
                "high",
                "validate",
                "Hexbin 联合分布图 + 边际直方图",
                "Python/matplotlib hexbin joint plot with marginal histograms",
                "large paired numeric samples can hide density under ordinary scatter points; hexbin exposes bivariate density and marginal distributions",
                "two continuous numeric sample columns with at least about 120 complete rows; report Pearson/Spearman correlation or quantiles nearby",
                hexbin_command(profile, x_col, y_col),
                visual_text,
            )
        )

    if polar_columns(profile) and 4 <= rows <= 40:
        ops.append(
            make_opportunity(
                profile,
                "polar",
                q,
                "medium",
                "result",
                "极坐标图/环形力分配图",
                "Python/matplotlib polar plot; Excel-style circular summary table as companion",
                "angle/direction data are naturally circular; Cartesian bars often hide periodic structure",
                "angle column plus magnitude/force/error column; keep exact values in a nearby table",
                polar_command_hint(profile),
                visual_text,
            )
        )

    if rows >= 200 and len(num) >= 1 and not has_bivariate_sample_cloud(profile):
        ops.append(
            make_opportunity(
                profile,
                "distribution",
                q,
                "high",
                "validate",
                "直方图/箱线图/小提琴图或概率分布图",
                "Python/matplotlib distribution plot; Excel histogram only for simple one-column checks",
                "large sample or Monte Carlo table should expose distribution, tail risk, and quantiles rather than only averages",
                "one or more numeric sample columns; report median/quantile/probability table next to the plot",
                distribution_command_hint(profile),
                visual_text,
            )
        )

    if ("sensitivity" in lower_path or "灵敏" in lower_path or "parameter_or_sensitivity" in profile.notes) and rows >= 5:
        ops.append(
            make_opportunity(
                profile,
                "sensitivity_curve",
                q,
                "medium",
                "validate",
                "灵敏度曲线/多情景对比图",
                "Python/matplotlib line or grouped bar; Excel line chart acceptable for simple tables",
                "sensitivity tables should reveal monotonicity, threshold, and stable interval",
                "parameter column plus response column; mark recommended value or threshold when available",
                line_command_hint(profile),
                visual_text,
            )
        )

    if any(term in lower_path for term in ["schedule", "timeline", "gantt", "process", "workflow", "实施", "流程"]):
        ops.append(
            make_opportunity(
                profile,
                "gantt",
                q,
                "medium",
                "operate",
                "甘特图/执行时间线",
                "Python/matplotlib Gantt; Excel timeline chart if dates/stages are simple",
                "workflow or schedule data are easier to review as an ordered timeline",
                "task/stage column plus start/end or duration column",
                gantt_command_hint(profile),
                visual_text,
            )
        )

    return ops


def contextual_opportunities(root: Path, visual_text: str) -> list[Opportunity]:
    context = collect_context(root)
    ops: list[Opportunity] = []
    if re.search(r"技术路线|建模思路|输入|输出|控制链|反演|状态|反馈|刚体|机制|流程", context):
        ops.append(
            context_opportunity(
                root,
                "global_ppt_reasoning",
                "ALL",
                "medium",
                "define",
                "PPT 可编辑技术路线图/机制示意图",
                "PowerPoint via plot_ppt_reasoning_diagram.py; export PNG/PDF for paper",
                "mechanism or route text is present; a PPT diagram can make model layers and dependencies visible",
                "diagram spec with nodes, lanes, arrows, and each shape bound to model variables/results",
                "python $PROJECT_ROOT\\\\.codex\\skills\\mira\\scripts\\plot_ppt_reasoning_diagram.py --root <project-root> --demo model_flow --prefix model_route",
                visual_text,
            )
        )
    if re.search(r"三维|圆柱|空间|层|通道|结构|架构|stack|tensor|channel", context, re.I):
        ops.append(
            context_opportunity(
                root,
                "global_pseudo_3d",
                "ALL",
                "optional",
                "define",
                "伪 3D PPT 结构示意图",
                "PowerPoint pseudo-3D diagram; keep PPTX editable and export PNG/PDF",
                "semantic depth, cylinder geometry, layer/channel, or scenario stack may be clearer as pseudo-3D schematic",
                "depth must represent a real variable/layer/scenario, not decoration",
                "python $PROJECT_ROOT\\\\.codex\\skills\\mira\\scripts\\plot_ppt_reasoning_diagram.py --root <project-root> --demo pseudo_3d_architecture --prefix pseudo3d_structure",
                visual_text,
            )
        )
    if re.search(r"算法|搜索|递推|动态规划|模拟退火|遗传|粒子群|蒙特卡洛|Monte Carlo|DP|SA|GA|PSO", context, re.I):
        ops.append(
            context_opportunity(
                root,
                "global_algorithm_flow",
                "ALL",
                "medium",
                "operate",
                "算法流程图/求解过程图",
                "PowerPoint flowchart or Python structure diagram",
                "nontrivial solver logic should be visible before code/results",
                "step labels must match solver states, updates, stopping rules, and output artifacts",
                "python $PROJECT_ROOT\\\\.codex\\skills\\mira\\scripts\\plot_ppt_reasoning_diagram.py --root <project-root> --demo decision_loop --prefix solver_flow",
                visual_text,
            )
        )
    return ops


def make_opportunity(
    profile: TableProfile,
    suffix: str,
    question: str,
    priority: str,
    role: str,
    visual: str,
    toolchain: str,
    reason: str,
    contract: str,
    command: str,
    visual_text: str,
) -> Opportunity:
    op_id = slug(f"{Path(profile.path).stem}-{suffix}")
    status, evidence = realization_status(visual_text, profile, visual, suffix)
    return Opportunity(op_id, question, profile.path, priority, role, visual, toolchain, reason, contract, command, status, evidence)


def context_opportunity(
    root: Path,
    op_id: str,
    question: str,
    priority: str,
    role: str,
    visual: str,
    toolchain: str,
    reason: str,
    contract: str,
    command: str,
    visual_text: str,
) -> Opportunity:
    status, evidence = realization_status_for_keywords(visual_text, visual_keywords(op_id, visual))
    return Opportunity(op_id, question, "planning/modeling/paper context", priority, role, visual, toolchain, reason, contract, command, status, evidence)


def realization_status(visual_text: str, profile: TableProfile, visual: str, suffix: str) -> tuple[str, str]:
    source_stem = Path(profile.path).stem.lower()
    keywords = visual_keywords(suffix, visual)
    text = visual_text.lower()
    exact = source_stem in text
    keyword = [term for term in keywords if term.lower() in text]
    question_hit = profile.question != "ALL" and profile.question.lower() in text and keyword
    if exact and keyword:
        return "realized", f"source stem and visual keyword found: {source_stem}, {keyword[0]}"
    if exact or question_hit or keyword:
        evidence = []
        if exact:
            evidence.append(f"source stem `{source_stem}` appears")
        if question_hit:
            evidence.append(f"{profile.question} visual keyword appears")
        if keyword:
            evidence.append("keyword " + ", ".join(keyword[:3]))
        return "partial", "; ".join(evidence)
    return "missing", "no matching visual/index/paper reference detected"


def realization_status_for_keywords(visual_text: str, keywords: list[str]) -> tuple[str, str]:
    text = visual_text.lower()
    hits = [term for term in keywords if term.lower() in text]
    if len(hits) >= 2:
        return "realized", "keywords " + ", ".join(hits[:4])
    if hits:
        return "partial", "keyword " + hits[0]
    return "missing", "no matching visual/index/paper reference detected"


def visual_keywords(kind: str, visual: str) -> list[str]:
    base = f"{kind} {visual}".lower()
    if "radar" in base or "雷达" in visual:
        return ["radar", "雷达"]
    if "parallel" in base or "平行" in visual:
        return ["parallel", "平行坐标"]
    if "heatmap" in base or "surface" in base or "热力" in visual or "曲面" in visual:
        return ["heatmap", "热力", "surface", "曲面", "contour", "等高线", "3d"]
    if "hexbin" in base or "joint" in base or "联合分布" in visual:
        return ["hexbin", "joint", "联合分布", "边际", "bivariate"]
    if "polar" in base or "极坐标" in visual:
        return ["polar", "极坐标", "rose", "角度", "方向"]
    if "distribution" in base or "分布" in visual or "小提琴" in visual:
        return ["distribution", "hist", "box", "violin", "分布", "残余", "monte"]
    if "sensitivity" in base or "灵敏" in visual:
        return ["sensitivity", "灵敏", "threshold", "阈值"]
    if "ppt" in base or "流程" in visual or "路线" in visual:
        return ["ppt", "flow", "流程", "路线", "framework", "chain", "control"]
    if "gantt" in base or "甘特" in visual:
        return ["gantt", "timeline", "甘特", "时间线"]
    return [kind.lower()]


def audit(
    root: Path,
    profiles: list[TableProfile],
    opportunities: list[Opportunity],
    three_d_candidates: list[ThreeDCandidate],
    visual_text: str,
    output_level: str,
) -> tuple[list[Finding], dict[str, Any]]:
    findings: list[Finding] = []
    contest_final = output_level.strip().lower() == "contest_final"
    high_missing = [op for op in opportunities if op.priority == "high" and op.status == "missing"]
    medium_missing = [op for op in opportunities if op.priority == "medium" and op.status == "missing"]
    radar_ops = [op for op in opportunities if "雷达" in op.recommended_visual or "radar" in op.command_hint.lower()]
    missing_radar = [op for op in radar_ops if op.status == "missing"]
    unresolved_three_d = [item for item in three_d_candidates if item.status == "unresolved"]
    figure_index_exists = (root / "figures" / "figure_index.md").exists() or (root / "diagrams" / "diagram_index.md").exists()

    for candidate in unresolved_three_d:
        findings.append(
            Finding(
                "FAIL" if contest_final and candidate.priority == "high" else "WARN",
                "three_d_opportunity_disposition",
                "implementation",
                f"unresolved 3D opportunity: {candidate.source} -> {candidate.kind}",
                "generate a truthful 3D artifact with a contour, projection, slice, or table companion, or record a specific evidence-backed waiver in planning/figure_storyboard.json.",
                f"{candidate.reason}; {candidate.disposition_issue}",
            )
        )

    if profiles and not figure_index_exists:
        findings.append(
            Finding(
                "WARN",
                "visual_index_missing",
                "implementation",
                "result tables exist but figure/diagram index is missing",
                "create figures/figure_index.md and diagrams/diagram_index.md with artifact, source data, script/toolchain, role, claim, and paper location.",
                f"{len(profiles)} table profiles scanned",
            )
        )

    for op in high_missing[:8]:
        findings.append(
            Finding(
                "WARN",
                "high_value_visual_opportunity",
                "implementation",
                f"missed high-priority visual opportunity: {op.source} -> {op.recommended_visual}",
                f"generate or explicitly waive this visual. Suggested toolchain: {op.recommended_toolchain}.",
                op.reason,
            )
        )

    if contest_final and len(medium_missing) >= 3:
        examples = "; ".join(f"{op.source}->{op.recommended_visual}" for op in medium_missing[:5])
        findings.append(
            Finding(
                "WARN",
                "medium_visual_opportunity_cluster",
                "implementation",
                f"{len(medium_missing)} medium-priority visual opportunities are missing",
                "select the highest-claim-value opportunities instead of repeating the same chart grammar.",
                examples,
            )
        )

    if missing_radar:
        examples = "; ".join(f"{op.source} ({op.priority})" for op in missing_radar[:5])
        findings.append(
            Finding(
                "WARN",
                "radar_chart_opportunity",
                "implementation",
                "multi-indicator table(s) could use radar charts or a normalized indicator visual, but no radar visual is detected",
                "when indicators can be normalized and directions aligned, use radar for small scheme sets; otherwise use parallel coordinates or a normalized score table.",
                examples,
            )
        )

    toolchains = Counter(op.recommended_toolchain.split(";")[0] for op in opportunities)
    missing_or_partial = [op for op in opportunities if op.status in {"missing", "partial"}]
    if len(opportunities) >= 5 and len({op.recommended_visual for op in opportunities}) >= 4 and not re.search(r"pptx|xlsx|matlab|octave|toolchain|工具链", visual_text, re.I):
        findings.append(
            Finding(
                "WARN",
                "visual_toolchain_not_recorded",
                "implementation",
                "visual opportunities span several grammars, but the final visual toolchain is not recorded in figure/diagram indexes",
                "record whether each figure was produced by Python, PPT/PowerPoint, Excel-compatible table/chart, MATLAB/Octave, or another tool. Do not claim a tool was used unless it actually ran.",
            )
        )

    metrics = {
        "table_profiles": len(profiles),
        "opportunities": len(opportunities),
        "opportunities_by_priority": dict(Counter(op.priority for op in opportunities)),
        "opportunities_by_status": dict(Counter(op.status for op in opportunities)),
        "opportunities_by_role": dict(Counter(op.role for op in opportunities)),
        "recommended_visuals": dict(Counter(op.recommended_visual for op in opportunities)),
        "recommended_toolchains": dict(toolchains),
        "high_missing": len(high_missing),
        "medium_missing": len(medium_missing),
        "missing_or_partial": len(missing_or_partial),
        "radar_opportunities": len(radar_ops),
        "missing_radar_opportunities": len(missing_radar),
        "hexbin_joint_opportunities": sum(1 for op in opportunities if "hexbin" in op.opportunity_id or "Hexbin" in op.recommended_visual),
        "three_d_candidates": len(three_d_candidates),
        "three_d_candidates_by_kind": dict(Counter(item.kind for item in three_d_candidates)),
        "three_d_candidates_by_status": dict(Counter(item.status for item in three_d_candidates)),
        "unresolved_three_d_candidates": len(unresolved_three_d),
        "figure_index_exists": figure_index_exists,
        "warnings": sum(1 for item in findings if item.level == "WARN"),
        "failures": sum(1 for item in findings if item.level == "FAIL"),
    }
    return findings, metrics


def is_metric_value_table(profile: TableProfile) -> bool:
    if profile.rows < 4 or profile.rows > 12 or len(profile.numeric_cols) != 1 or len(profile.text_cols) != 1:
        return False
    headers = " ".join(profile.headers).lower()
    return bool(re.search(r"metric|指标|value|得分|score|性能|策略|strategy", headers + " " + profile.path.lower()))


def has_scheme_metric_matrix(profile: TableProfile) -> bool:
    if not (2 <= profile.rows <= 30 and 3 <= len(profile.numeric_cols) <= 10):
        return False
    if not profile.text_cols:
        return False
    headers = " ".join(profile.headers).lower()
    path = profile.path.lower()
    return bool(re.search(r"scheme|method|model|case|scenario|方案|方法|模型|策略|候选|player|队员|allocation|ranking|score", headers + " " + path))


def has_two_parameter_response(profile: TableProfile) -> bool:
    if profile.rows < 9 or len(profile.numeric_cols) < 3:
        return False
    headers = " ".join(profile.headers).lower()
    path = profile.path.lower()
    sensitivity = re.search(r"noise|sd|sigma|duration|参数|灵敏|sensitivity|alpha|lambda|cost|权重|grid|surface", headers + " " + path)
    response = re.search(r"prob|概率|rate|success|成功|residual|error|objective|目标|cost|value|score|median|p90", headers)
    return bool(sensitivity and response)


def has_bivariate_sample_cloud(profile: TableProfile) -> bool:
    if profile.rows < 120 or len(profile.numeric_cols) < 2:
        return False
    headers = " ".join(profile.headers).lower()
    path = profile.path.lower()
    if re.search(r"grid|surface|matrix|pivot|矩阵|网格|曲面|热力|heatmap", headers + " " + path):
        return False
    if len(profile.text_cols) > 3:
        return False
    sample_terms = re.search(
        r"sample|monte|simulation|residual|error|pred|obs|observed|feature|response|"
        r"样本|仿真|模拟|蒙特|残差|误差|预测|观测|特征|响应|散点|联合",
        headers + " " + path,
        re.I,
    )
    return bool(sample_terms or (profile.rows >= 300 and len(profile.numeric_cols) <= 6))


def bivariate_sample_columns(profile: TableProfile) -> tuple[str, str]:
    preferred = [
        col
        for col in profile.numeric_cols
        if not re.search(r"^id$|编号|序号|index|rank|time_step|iteration|迭代", col, re.I)
    ]
    cols = preferred or profile.numeric_cols
    if len(cols) >= 2:
        return cols[0], cols[1]
    return profile.numeric_cols[0], profile.numeric_cols[0]


def response_columns(profile: TableProfile) -> list[str]:
    return [
        col
        for col in profile.numeric_cols
        if re.search(r"prob|概率|rate|success|成功|residual|error|objective|目标|cost|value|score|median|p90", col, re.I)
    ]


def parameter_columns(profile: TableProfile) -> list[str]:
    preferred = []
    fallback = []
    for col in profile.numeric_cols:
        if re.search(r"noise|sd|sigma|duration|参数|alpha|lambda|weight|权重|force_noise|timing_noise|delta_|脉冲", col, re.I):
            preferred.append(col)
        elif not re.search(r"residual|error|angle|prob|概率|rate|success|成功|objective|目标|cost|value|score|median|p90", col, re.I):
            fallback.append(col)
    return preferred + fallback


def radar_command(profile: TableProfile) -> str:
    if is_metric_value_table(profile):
        return (
            "reshape metric-value table into scheme x normalized metrics if needed, then run "
            "plot_radar_taylor_surface.py --kind radar --input <normalized-table> --item-col 方案 --metrics <指标列> --normalize minmax"
        )
    item_col = profile.text_cols[0] if profile.text_cols else "方案"
    metrics = ",".join(profile.numeric_cols[:6])
    return (
        "python $PROJECT_ROOT\\\\.codex\\skills\\mira\\scripts\\plot_radar_taylor_surface.py "
        f"--root <project-root> --kind radar --input {profile.path} --item-col {item_col} --metrics {metrics} --normalize minmax --prefix {Path(profile.path).stem}_radar"
    )


def parallel_command(profile: TableProfile) -> str:
    item_col = profile.text_cols[0] if profile.text_cols else "方案"
    metrics = ",".join(profile.numeric_cols[:8])
    return (
        "python $PROJECT_ROOT\\\\.codex\\skills\\mira\\scripts\\plot_parallel_pareto_gantt.py "
        f"--root <project-root> --kind parallel --input {profile.path} --item-col {item_col} --metrics {metrics} --normalize minmax --prefix {Path(profile.path).stem}_parallel"
    )


def surface_command(profile: TableProfile, response: str) -> str:
    params = parameter_columns(profile)
    if len(params) >= 2:
        x, y = params[:2]
    else:
        x, y = profile.numeric_cols[:2]
    z = response or profile.numeric_cols[-1]
    return (
        "python $PROJECT_ROOT\\\\.codex\\skills\\mira\\scripts\\plot_radar_taylor_surface.py "
        f"--root <project-root> --kind surface --input {profile.path} --x {x} --y {y} --z {z} --contour --prefix {Path(profile.path).stem}_surface"
    )


def polar_command_hint(profile: TableProfile) -> str:
    cols = polar_columns(profile)
    if cols:
        angle, magnitude = cols
    else:
        angle, magnitude = "angle", "magnitude"
    return f"use matplotlib polar axes with angle `{angle}` and magnitude `{magnitude}`; keep exact values in a table"


def polar_columns(profile: TableProfile) -> tuple[str, str] | None:
    """Return a credible angular coordinate and radial value for polar plots.

    Do not treat response angles such as tilt_angle/residual_angle as a polar
    coordinate by themselves; those are usually y-values for sensitivity curves.
    """

    angle_candidates = [
        col
        for col in profile.numeric_cols
        if re.search(r"direction|theta|azimuth|bearing|方位|方向|angle_deg|角度", col, re.I)
        and not re.search(r"tilt_angle|residual_angle|倾角|偏角|median|p90|prob|概率", col, re.I)
    ]
    if not angle_candidates:
        return None
    angle = next((col for col in angle_candidates if re.search(r"direction|方向|theta|azimuth|bearing", col, re.I)), angle_candidates[0])
    magnitude_candidates = [
        col
        for col in profile.numeric_cols
        if col != angle
        and not re.search(r"^case$|^player$|编号|序号|index|duration|time|prob|概率|rate|success|成功", col, re.I)
    ]
    preferred = [
        col
        for col in magnitude_candidates
        if re.search(r"force|magnitude|value|tilt_angle|residual_angle|delta|幅值|力|倾角|偏角|误差", col, re.I)
    ]
    magnitude = (preferred or magnitude_candidates or [None])[0]
    if not magnitude:
        return None
    return angle, magnitude


def distribution_command_hint(profile: TableProfile) -> str:
    value = profile.numeric_cols[0] if profile.numeric_cols else "value"
    return f"python plot_distribution.py --root <project-root> --input {profile.path} --value-cols {value} --prefix {Path(profile.path).stem}_distribution"


def hexbin_command(profile: TableProfile, x_col: str, y_col: str) -> str:
    return (
        "python $PROJECT_ROOT\\\\.codex\\skills\\mira\\scripts\\plot_hexbin_joint.py "
        f"--root <project-root> --input {profile.path} --x {x_col} --y {y_col} --prefix {Path(profile.path).stem}_hexbin"
    )


def line_command_hint(profile: TableProfile) -> str:
    x = profile.numeric_cols[0] if profile.numeric_cols else profile.headers[0]
    ys = ",".join(profile.numeric_cols[1:4]) if len(profile.numeric_cols) > 1 else "<response-col>"
    return f"use Python/Excel line chart with x `{x}` and response columns `{ys}`; mark threshold/recommended parameter if available"


def gantt_command_hint(profile: TableProfile) -> str:
    return f"python plot_parallel_pareto_gantt.py --root <project-root> --kind gantt --input {profile.path} --task-col <task> --start-col <start> --end-col <end>"


def collect_visual_text(root: Path) -> str:
    parts: list[str] = []
    for rel_path in ["figures/figure_index.md", "diagrams/diagram_index.md", "planning/figure_storyboard.md", "paper/main.tex"]:
        parts.append(read_text(root / rel_path))
    for candidate in PAPER_CANDIDATES:
        parts.append(read_text(root / candidate))
    for base in [root / "figures", root / "diagrams"]:
        if base.exists():
            for path in sorted(base.rglob("*")):
                if path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".svg", ".pdf", ".pptx", ".xlsx"}:
                    parts.append(rel(root, path))
    return "\n".join(parts)


def collect_context(root: Path) -> str:
    rels = [
        "planning/problem_analysis.md",
        "planning/modeling_plan.md",
        "planning/method_route.md",
        "planning/modeling_route.md",
        "planning/validation_plan.md",
        "planning/evidence_plan.md",
        "results/result_report.md",
        "planning/result_ledger.md",
        "paper/main.tex",
    ]
    return "\n".join(read_text(root / rel_path)[:20000] for rel_path in rels)


def detect_question(name: str) -> str:
    match = QUESTION_RE.search(name)
    if not match:
        return "ALL"
    if match.group(1):
        return f"Q{int(match.group(1))}"
    return f"Q{CN_NUMBERS.get(match.group(2), 0)}"


def dedupe_opportunities(opportunities: list[Opportunity]) -> list[Opportunity]:
    seen: set[str] = set()
    out: list[Opportunity] = []
    for op in opportunities:
        key = f"{op.opportunity_id}|{op.source}|{op.recommended_visual}"
        if key in seen:
            continue
        seen.add(key)
        out.append(op)
    return out


def is_number(value: Any) -> bool:
    try:
        float(str(value).strip())
        return True
    except ValueError:
        return False


def verdict(findings: list[Finding]) -> str:
    if any(item.level == "FAIL" for item in findings):
        return "FAIL"
    if any(item.level == "WARN" for item in findings):
        return "PASS_WITH_WARNINGS"
    return "PASS"


def markdown(payload: dict[str, Any]) -> str:
    lines = [
        "# Mira 0.8.2 Visual Opportunity Audit",
        "",
        f"- Generated: {payload['generated_at']}",
        f"- Verdict: **{payload['verdict']}**",
        f"- Root: `{payload['root']}`",
        "",
        "## Metrics",
        "",
        "| Metric | Value |",
        "|---|---|",
    ]
    for key, value in payload["metrics"].items():
        lines.append(f"| {key} | {escape(json.dumps(value, ensure_ascii=False))} |")

    lines.extend(["", "## Opportunities", "", "| ID | Priority | Status | Question | Source | Role | Recommended visual | Toolchain | Reason | Command hint |", "|---|---|---|---|---|---|---|---|---|---|"])
    for op in payload["opportunities"]:
        lines.append(
            "| {id} | {priority} | {status} | {question} | `{source}` | `{role}` | {visual} | {toolchain} | {reason} | `{command}` |".format(
                id=escape(op["opportunity_id"]),
                priority=op["priority"],
                status=op["status"],
                question=op["question"],
                source=escape(op["source"]),
                role=op["role"],
                visual=escape(op["recommended_visual"]),
                toolchain=escape(op["recommended_toolchain"]),
                reason=escape(op["reason"]),
                command=escape(op["command_hint"]),
            )
        )

    lines.extend(
        [
            "",
            "## Three-Dimensional Candidates",
            "",
            "| ID | Kind | Status | Question | Source | Data columns | Required companion | Reason | Disposition issue |",
            "|---|---|---|---|---|---|---|---|---|",
        ]
    )
    if payload["three_d_candidates"]:
        for item in payload["three_d_candidates"]:
            lines.append(
                f"| `{escape(item['candidate_id'])}` | `{escape(item['kind'])}` | `{escape(item['status'])}` | "
                f"{escape(item['question'])} | `{escape(item['source'])}` | "
                f"{escape(', '.join(item['data_columns']))} | `{escape(item['required_companion'])}` | "
                f"{escape(item['reason'])} | {escape(item['disposition_issue'])} |"
            )
    else:
        lines.append("| - | - | - | - | - | - | - | no data-driven 3D opportunity detected | - |")

    lines.extend(["", "## Findings", "", "| Level | Axis | Return to | Finding | Recommendation | Evidence |", "|---|---|---|---|---|---|"])
    if payload["findings"]:
        for item in payload["findings"]:
            lines.append(
                f"| {item['level']} | {item['axis']} | {item['return_phase']} | {escape(item['finding'])} | {escape(item['recommendation'])} | {escape(item.get('evidence', ''))} |"
            )
    else:
        lines.append("| PASS | visual_opportunity | - | no missed high-value visual opportunity detected | keep toolchain recorded in figure index | - |")
    lines.append("")
    return "\n".join(lines)


def emit(payload: dict[str, Any]) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    print(f"VERDICT: {payload['verdict']}")
    print("metrics: " + json.dumps(payload["metrics"], ensure_ascii=False, sort_keys=True))
    for item in payload["findings"]:
        print(f"{item['level']}: {item['axis']}: {item['finding']}")


def read_text(path: Path) -> str:
    if not path.exists():
        return ""
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def load_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(read_text(path))
        return value if isinstance(value, dict) else {}
    except json.JSONDecodeError:
        return {}


def clean_list(value: Any) -> list[str]:
    return [str(item).strip() for item in value if str(item).strip()] if isinstance(value, list) else []


def is_placeholder(value: Any) -> bool:
    return str(value or "").strip().lower() in {"", "tbd", "todo", "pending", "unresolved", "none", "n/a"}


def resolve_path(root: Path, value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else root / path


def rel(root: Path, path: Path) -> str:
    try:
        return str(path.resolve().relative_to(root)).replace("\\", "/")
    except ValueError:
        return str(path)


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "visual-opportunity"


def escape(value: Any) -> str:
    return str(value).replace("|", "\\|").replace("`", "'").replace("\n", " ")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=".", help="Contest project root")
    parser.add_argument("--output-level", default="contest_final", help="quick_draft, reproducible_draft, or contest_final")
    parser.add_argument("--write-report", help="Write markdown report")
    parser.add_argument("--write-json", help="Write JSON report")
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(main())
