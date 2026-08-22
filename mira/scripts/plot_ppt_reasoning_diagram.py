#!/usr/bin/env python3
"""Generate editable PowerPoint reasoning diagrams for Mira papers."""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
except Exception as exc:  # pragma: no cover - dependency guard
    raise SystemExit("ERROR: python-pptx is required. Install package `python-pptx`.") from exc


SLIDE_W = 13.333
SLIDE_H = 7.5
CANVAS = {"left": 0.55, "top": 0.92, "right": 0.45, "bottom": 0.38}

COLORS = {
    "blue": "#2F6B9A",
    "teal": "#2A9D8F",
    "orange": "#E76F51",
    "gold": "#E9C46A",
    "green": "#5E8C61",
    "purple": "#7B5EA7",
    "red": "#B23A48",
    "gray": "#6B7280",
    "line": "#2F4A6D",
    "text": "#1F2933",
    "lane": "#EEF4FA",
    "white": "#FFFFFF",
}

NODE_STYLES = {
    "data": {"fill": "#E8F2FB", "line": COLORS["blue"], "shape": "round"},
    "process": {"fill": "#F4EFE8", "line": COLORS["orange"], "shape": "round"},
    "model": {"fill": "#EAF4EF", "line": COLORS["teal"], "shape": "round"},
    "decision": {"fill": "#FFF6D8", "line": COLORS["gold"], "shape": "diamond"},
    "state": {"fill": "#F1ECF7", "line": COLORS["purple"], "shape": "hex"},
    "equation": {"fill": "#FFFFFF", "line": COLORS["gray"], "shape": "rect"},
    "result": {"fill": "#E9F5E6", "line": COLORS["green"], "shape": "round"},
    "warning": {"fill": "#FCEEEF", "line": COLORS["red"], "shape": "round"},
    "title": {"fill": "#EAF4EF", "line": COLORS["teal"], "shape": "round"},
    "default": {"fill": "#FFFFFF", "line": COLORS["line"], "shape": "round"},
}

SHAPES = {
    "round": MSO_SHAPE.ROUNDED_RECTANGLE,
    "rounded": MSO_SHAPE.ROUNDED_RECTANGLE,
    "rect": MSO_SHAPE.RECTANGLE,
    "rectangle": MSO_SHAPE.RECTANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "ellipse": MSO_SHAPE.OVAL,
    "oval": MSO_SHAPE.OVAL,
    "hex": MSO_SHAPE.HEXAGON,
    "hexagon": MSO_SHAPE.HEXAGON,
    "parallelogram": MSO_SHAPE.PARALLELOGRAM,
    "chevron": MSO_SHAPE.CHEVRON,
}


def main() -> int:
    args = parse_args()
    root = Path(args.root).resolve()
    spec, source = load_spec(args, root)
    spec = normalize_spec(spec, args)
    outputs = prepare_outputs(args, root)

    write_csv(outputs["nodes"], spec.get("nodes", []))
    write_csv(outputs["edges"], spec.get("edges", []))
    write_csv(outputs["groups"], spec.get("groups", []))
    write_csv(outputs["callouts"], spec.get("callouts", []))
    outputs["spec"].write_text(json.dumps(spec, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    pptx_path = outputs["pptx"]
    build_pptx(spec, pptx_path, args)

    exported, export_notes = export_outputs(
        pptx_path,
        outputs["png"],
        outputs["pdf"],
        args.export,
        args.png_width,
        args.png_height,
    )
    if args.require_export and not exported:
        raise SystemExit("ERROR: export was required but no PNG/PDF was produced: " + "; ".join(export_notes))

    written = [str(pptx_path)] + [str(path) for path in [outputs["png"], outputs["pdf"]] if path.exists()]
    params = {
        "source": source,
        "export": args.export,
        "exported": exported,
        "export_notes": export_notes,
        "slide_size_inches": [SLIDE_W, SLIDE_H],
        "n_nodes": len(spec.get("nodes", [])),
        "n_edges": len(spec.get("edges", [])),
        "n_groups": len(spec.get("groups", [])),
        "n_callouts": len(spec.get("callouts", [])),
        "written": written,
    }
    outputs["params"].write_text(json.dumps(params, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_index(outputs, args, source, spec, written, export_notes)

    print("INFO: PPT reasoning diagram written:")
    print(f"  pptx: {pptx_path}")
    for path in [outputs["png"], outputs["pdf"]]:
        if path.exists():
            print(f"  export: {path}")
    print(f"  spec: {outputs['spec']}")
    print(f"  nodes: {outputs['nodes']}")
    print(f"  edges: {outputs['edges']}")
    print(f"  params: {outputs['params']}")
    print(f"  index: {outputs['index']}")
    if export_notes:
        print("  export_notes: " + " | ".join(export_notes))
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=".", help="Project root for relative paths.")
    parser.add_argument("--spec", help="JSON diagram specification.")
    parser.add_argument("--demo", choices=("model_flow", "decision_loop", "paper_logic", "showcase_work", "multimodal_fusion", "cnn_architecture", "pseudo_3d_architecture"), help="Generate a built-in demo.")
    parser.add_argument("--orientation", choices=("horizontal", "vertical"), default="horizontal", help="Auto-layout direction.")
    parser.add_argument("--diagrams-dir", "--out-dir", dest="diagrams_dir", default="diagrams", help="Diagram output directory.")
    parser.add_argument("--data-dir", default="results/figures_data", help="Source-data output directory.")
    parser.add_argument("--prefix", default="ppt_reasoning", help="Output filename prefix.")
    parser.add_argument("--title", default="", help="Optional title override.")
    parser.add_argument("--export", choices=("auto", "none", "office", "libreoffice"), default="auto", help="Export PPTX to PNG/PDF when possible.")
    parser.add_argument("--require-export", action="store_true", help="Fail if PNG/PDF export is unavailable.")
    parser.add_argument("--png-width", type=int, default=2400, help="PowerPoint PNG export width.")
    parser.add_argument("--png-height", type=int, default=1350, help="PowerPoint PNG export height.")
    return parser.parse_args()


def load_spec(args: argparse.Namespace, root: Path) -> tuple[dict[str, Any], str]:
    if args.demo:
        return demo_spec(args.demo), f"demo:{args.demo}"
    if args.spec:
        path = resolve_path(root, args.spec)
        if not path.exists():
            raise SystemExit(f"ERROR: spec file does not exist: {path}")
        return json.loads(path.read_text(encoding="utf-8")), str(path)
    raise SystemExit("ERROR: provide --spec <json> or --demo <model_flow|decision_loop|paper_logic|showcase_work|multimodal_fusion|cnn_architecture|pseudo_3d_architecture>.")


def normalize_spec(spec: dict[str, Any], args: argparse.Namespace) -> dict[str, Any]:
    spec = json.loads(json.dumps(spec, ensure_ascii=False))
    if args.title:
        spec["title"] = args.title
    spec.setdefault("title", "PPT 思路图")
    spec.setdefault("caption", "图：PPT 形状图展示模型推理链条与执行关系。")
    spec.setdefault("nodes", [])
    spec.setdefault("edges", [])
    spec.setdefault("groups", [])
    spec.setdefault("callouts", [])
    if not spec["nodes"]:
        raise SystemExit("ERROR: diagram spec requires at least one node.")

    seen: set[str] = set()
    for index, node in enumerate(spec["nodes"]):
        node.setdefault("id", f"n{index + 1}")
        node["id"] = str(node["id"])
        if node["id"] in seen:
            raise SystemExit(f"ERROR: duplicate node id: {node['id']}")
        seen.add(node["id"])
        node.setdefault("label", node["id"])
        node.setdefault("type", "default")

    valid_ids = {node["id"] for node in spec["nodes"]}
    spec["edges"] = [
        edge for edge in spec["edges"]
        if str(edge.get("source", "")) in valid_ids and str(edge.get("target", "")) in valid_ids
    ]
    if not all("x" in node and "y" in node for node in spec["nodes"]):
        auto_layout(spec, args.orientation)
    normalize_geometry(spec)
    return spec


def auto_layout(spec: dict[str, Any], orientation: str) -> None:
    nodes = spec["nodes"]
    ids = [node["id"] for node in nodes]
    order = {node_id: i for i, node_id in enumerate(ids)}
    adj: dict[str, list[str]] = {node_id: [] for node_id in ids}
    indeg: dict[str, int] = {node_id: 0 for node_id in ids}
    for edge in spec.get("edges", []):
        source, target = str(edge["source"]), str(edge["target"])
        adj[source].append(target)
        indeg[target] += 1

    queue = sorted([node_id for node_id, deg in indeg.items() if deg == 0], key=order.get)
    layer = {node_id: 0 for node_id in ids}
    visited: list[str] = []
    while queue:
        node_id = queue.pop(0)
        visited.append(node_id)
        for target in adj[node_id]:
            layer[target] = max(layer[target], layer[node_id] + 1)
            indeg[target] -= 1
            if indeg[target] == 0:
                queue.append(target)
                queue.sort(key=order.get)
    for node_id in ids:
        if node_id not in visited:
            layer[node_id] = max(layer.values() or [0]) + 1

    layer_values = sorted(set(layer.values()))
    layer_index = {value: i for i, value in enumerate(layer_values)}
    buckets: dict[int, list[dict[str, Any]]] = {layer_index[value]: [] for value in layer_values}
    for node in nodes:
        buckets[layer_index[layer[node["id"]]]].append(node)

    n_layers = max(1, len(buckets))
    for li, bucket in buckets.items():
        bucket.sort(key=lambda item: order[item["id"]])
        for bi, node in enumerate(bucket):
            along = 0.5 if n_layers == 1 else 0.06 + 0.88 * li / (n_layers - 1)
            cross = (bi + 1) / (len(bucket) + 1)
            if orientation == "vertical":
                node["x"], node["y"] = cross, along
            else:
                node["x"], node["y"] = along, cross


def normalize_geometry(spec: dict[str, Any]) -> None:
    for node in spec["nodes"]:
        node["x"] = clamp(float(node.get("x", 0.5)), 0.03, 0.97)
        node["y"] = clamp(float(node.get("y", 0.5)), 0.05, 0.95)
        node["w"] = clamp(float(node.get("w", 0.14)), 0.07, 0.26)
        node["h"] = clamp(float(node.get("h", 0.12)), 0.06, 0.20)
    for group in spec.get("groups", []):
        group["x"] = clamp(float(group.get("x", 0.05)), 0.0, 0.98)
        group["y"] = clamp(float(group.get("y", 0.08)), 0.0, 0.98)
        group["w"] = clamp(float(group.get("w", 0.25)), 0.05, 1.0 - group["x"])
        group["h"] = clamp(float(group.get("h", 0.80)), 0.05, 1.0 - group["y"])
    node_map = {node["id"]: node for node in spec["nodes"]}
    for callout in spec.get("callouts", []):
        target = node_map.get(str(callout.get("target", "")))
        if "x" not in callout:
            callout["x"] = (target["x"] + 0.10) if target else 0.72
        if "y" not in callout:
            callout["y"] = (target["y"] - 0.14) if target else 0.18
        callout["x"] = clamp(float(callout["x"]), 0.03, 0.82)
        callout["y"] = clamp(float(callout["y"]), 0.03, 0.88)
        callout["w"] = clamp(float(callout.get("w", 0.20)), 0.10, 0.33)
        callout["h"] = clamp(float(callout.get("h", 0.11)), 0.06, 0.20)


def build_pptx(spec: dict[str, Any], pptx_path: Path, args: argparse.Namespace) -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_title(slide, spec["title"])
    for group in spec.get("groups", []):
        add_group(slide, group)

    nodes = {node["id"]: node for node in spec["nodes"]}
    for edge in spec.get("edges", []):
        add_edge(slide, nodes[str(edge["source"])], nodes[str(edge["target"])], edge, args.orientation)
    for node in spec["nodes"]:
        add_node(slide, node)
    for callout in spec.get("callouts", []):
        add_callout(slide, callout, nodes)

    add_footer(slide, spec)
    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(pptx_path)


def add_title(slide: Any, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.18), Inches(12.3), Inches(0.42))
    set_text(box, title, size=22, bold=True, align=PP_ALIGN.CENTER)


def add_footer(slide: Any, spec: dict[str, Any]) -> None:
    note = spec.get("note", "")
    if not note:
        return
    box = slide.shapes.add_textbox(Inches(0.70), Inches(7.03), Inches(11.9), Inches(0.22))
    set_text(box, note, size=8, color=COLORS["gray"], align=PP_ALIGN.CENTER)


def add_group(slide: Any, group: dict[str, Any]) -> None:
    left, top, width, height = to_box(group["x"], group["y"], group["w"], group["h"])
    shape_type = SHAPES.get(str(group.get("shape", "round")).lower(), MSO_SHAPE.ROUNDED_RECTANGLE)
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    fill = group.get("fill") or style_for(group.get("type", "default"))["fill"]
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.fill.transparency = int(group.get("transparency", 35))
    shape.line.color.rgb = rgb(group.get("line", "#CBD5E1"))
    shape.line.width = Pt(float(group.get("line_width", 0.6)))
    if bool(group.get("dashed", False)):
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    label = str(group.get("label", ""))
    if label:
        label_box = slide.shapes.add_textbox(left + Inches(0.13), top + Inches(0.10), min(width, Inches(1.7)), Inches(0.25))
        set_text(label_box, label, size=int(group.get("font_size", 10)), bold=True, color=group.get("text_color", COLORS["gray"]), align=PP_ALIGN.LEFT)


def add_node(slide: Any, node: dict[str, Any]) -> None:
    style = style_for(node.get("type", "default"))
    left, top, width, height = to_box(node["x"] - node["w"] / 2, node["y"] - node["h"] / 2, node["w"], node["h"])
    shape_type = SHAPES.get(str(node.get("shape") or style["shape"]).lower(), MSO_SHAPE.ROUNDED_RECTANGLE)
    add_pseudo_3d_layers(slide, node, style, shape_type)
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(node.get("fill", style["fill"]))
    shape.fill.transparency = int(node.get("transparency", 0))
    shape.line.color.rgb = rgb(node.get("line", style["line"]))
    shape.line.width = Pt(float(node.get("line_width", 1.2)))
    if bool(node.get("dashed", False)):
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    set_text(
        shape,
        str(node.get("label", node["id"])),
        size=int(node.get("font_size", 10)),
        bold=bool(node.get("bold", False)),
        color=node.get("text_color", COLORS["text"]),
    )


def add_pseudo_3d_layers(slide: Any, node: dict[str, Any], style: dict[str, str], shape_type: Any) -> None:
    if not bool(node.get("pseudo_3d", False)) and int(node.get("depth_layers", 1)) <= 1:
        return
    layers = clamp_int(int(node.get("depth_layers", 3)), 1, 8)
    dx = float(node.get("depth_dx", 0.012))
    dy = float(node.get("depth_dy", -0.018))
    fill = node.get("depth_fill", node.get("fill", style["fill"]))
    line = node.get("depth_line", node.get("line", style["line"]))
    transparency = int(node.get("depth_transparency", 45))
    for step in range(layers, 0, -1):
        left, top, width, height = to_box(
            node["x"] - node["w"] / 2 + dx * step,
            node["y"] - node["h"] / 2 + dy * step,
            node["w"],
            node["h"],
        )
        layer_shape = slide.shapes.add_shape(shape_type, left, top, width, height)
        layer_shape.fill.solid()
        layer_shape.fill.fore_color.rgb = rgb(fill)
        layer_shape.fill.transparency = transparency
        layer_shape.line.color.rgb = rgb(line)
        layer_shape.line.width = Pt(float(node.get("depth_line_width", 0.7)))


def add_edge(slide: Any, source: dict[str, Any], target: dict[str, Any], edge: dict[str, Any], orientation: str) -> None:
    if str(edge.get("style", "")).lower() == "feedback":
        add_feedback_edge(slide, source, target, edge)
        return
    x1, y1, x2, y2 = edge_points(source, target, orientation)
    connector_type = MSO_CONNECTOR.ELBOW if str(edge.get("connector", "")).lower() == "elbow" else MSO_CONNECTOR.STRAIGHT
    connector = slide.shapes.add_connector(connector_type, *to_point(x1, y1), *to_point(x2, y2))
    color = edge_color(edge)
    connector.line.color.rgb = rgb(edge.get("color", color))
    connector.line.width = Pt(float(edge.get("width", 1.35)))
    if bool(edge.get("dashed", False)) or str(edge.get("style", "")).lower() in {"feedback", "assumption", "optional"}:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    set_line_arrow(connector)

    label = str(edge.get("label", "")).strip()
    if label:
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        left, top, width, height = to_box(clamp(mx - 0.055, 0, 0.90), clamp(my - 0.035, 0, 0.95), 0.11, 0.055)
        box = slide.shapes.add_textbox(left, top, width, height)
        set_text(box, label, size=int(edge.get("font_size", 8)), color=edge.get("color", color), align=PP_ALIGN.CENTER)


def add_feedback_edge(slide: Any, source: dict[str, Any], target: dict[str, Any], edge: dict[str, Any]) -> None:
    color = edge.get("color", COLORS["red"])
    sx, sy = source["x"], source["y"] - source["h"] / 2
    tx, ty = target["x"], target["y"] - target["h"] / 2
    route_y = clamp(min(sy, ty) - float(edge.get("lift", 0.16)), 0.03, 0.92)
    points = [(sx, sy), (sx, route_y), (tx, route_y), (tx, ty)]
    for index, (start, end) in enumerate(zip(points, points[1:])):
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *to_point(*start), *to_point(*end))
        connector.line.color.rgb = rgb(color)
        connector.line.width = Pt(float(edge.get("width", 1.15)))
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if index == len(points) - 2:
            set_line_arrow(connector)
    label = str(edge.get("label", "")).strip()
    if label:
        mx = (sx + tx) / 2
        left, top, width, height = to_box(clamp(mx - 0.06, 0, 0.88), clamp(route_y - 0.045, 0, 0.94), 0.12, 0.055)
        box = slide.shapes.add_textbox(left, top, width, height)
        set_text(box, label, size=int(edge.get("font_size", 8)), color=color, align=PP_ALIGN.CENTER)


def add_callout(slide: Any, callout: dict[str, Any], nodes: dict[str, dict[str, Any]]) -> None:
    left, top, width, height = to_box(callout["x"], callout["y"], callout["w"], callout["h"])
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = rgb(callout.get("fill", "#FFF8E7"))
    box.line.color.rgb = rgb(callout.get("line", COLORS["gold"]))
    box.line.width = Pt(1.0)
    set_text(box, str(callout.get("text", "")), size=int(callout.get("font_size", 8)), color=callout.get("text_color", COLORS["text"]), align=PP_ALIGN.LEFT)

    target = nodes.get(str(callout.get("target", "")))
    if target:
        x1, y1 = callout["x"], callout["y"] + callout["h"] / 2
        x2, y2 = target["x"], target["y"]
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *to_point(x1, y1), *to_point(x2, y2))
        connector.line.color.rgb = rgb(callout.get("line", COLORS["gold"]))
        connector.line.width = Pt(0.9)
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def edge_points(source: dict[str, Any], target: dict[str, Any], orientation: str) -> tuple[float, float, float, float]:
    sx, sy, sw, sh = source["x"], source["y"], source["w"], source["h"]
    tx, ty, tw, th = target["x"], target["y"], target["w"], target["h"]
    dx, dy = tx - sx, ty - sy
    if orientation == "vertical" or abs(dy) > abs(dx):
        y1 = sy + (sh / 2 if dy >= 0 else -sh / 2)
        y2 = ty - (th / 2 if dy >= 0 else -th / 2)
        return sx, y1, tx, y2
    x1 = sx + (sw / 2 if dx >= 0 else -sw / 2)
    x2 = tx - (tw / 2 if dx >= 0 else -tw / 2)
    return x1, sy, x2, ty


def set_text(shape: Any, text: str, *, size: int, bold: bool = False, color: str = COLORS["text"], align: Any = PP_ALIGN.CENTER) -> None:
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text_frame.margin_left = Inches(0.05)
    text_frame.margin_right = Inches(0.05)
    text_frame.margin_top = Inches(0.03)
    text_frame.margin_bottom = Inches(0.03)
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def set_line_arrow(connector: Any) -> None:
    line = connector._element.spPr.get_or_add_ln()
    existing = line.find(qn("a:tailEnd"))
    if existing is not None:
        line.remove(existing)
    arrow = OxmlElement("a:tailEnd")
    arrow.set("type", "triangle")
    arrow.set("w", "sm")
    arrow.set("len", "sm")
    line.append(arrow)


def export_outputs(pptx_path: Path, png_path: Path, pdf_path: Path, mode: str, width: int, height: int) -> tuple[bool, list[str]]:
    if mode == "none":
        return False, ["export skipped by --export none"]
    notes: list[str] = []
    if mode in {"auto", "office"}:
        ok, note = export_with_powerpoint(pptx_path, png_path, pdf_path, width, height)
        notes.append(note)
        if ok:
            return True, notes
        if mode == "office":
            return False, notes
    if mode in {"auto", "libreoffice"}:
        ok, note = export_with_libreoffice(pptx_path, pdf_path)
        notes.append(note)
        if ok:
            return True, notes
        if mode == "libreoffice":
            return False, notes
    notes.append("no export backend available; PPTX source was still written")
    return False, notes


def export_with_powerpoint(pptx_path: Path, png_path: Path, pdf_path: Path, width: int, height: int) -> tuple[bool, str]:
    script = f"""
$ErrorActionPreference = 'Stop'
$pptx = {ps_quote(pptx_path)}
$png = {ps_quote(png_path)}
$pdf = {ps_quote(pdf_path)}
$app = $null
$pres = $null
try {{
  $app = New-Object -ComObject PowerPoint.Application
  $pres = $app.Presentations.Open($pptx, $true, $false, $false)
  $pres.SaveAs($pdf, 32)
  $pres.Slides.Item(1).Export($png, 'PNG', {int(width)}, {int(height)})
  'EXPORT_OK'
}} finally {{
  if ($pres -ne $null) {{ $pres.Close() }}
  if ($app -ne $null) {{ $app.Quit() }}
}}
"""
    try:
        result = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
            capture_output=True,
            text=True,
            timeout=120,
        )
    except Exception as exc:
        return False, f"PowerPoint COM export failed to start: {exc}"
    if result.returncode == 0 and png_path.exists() and pdf_path.exists():
        return True, "exported via PowerPoint COM"
    detail = (result.stderr or result.stdout or "").strip().replace("\n", " ")
    return False, "PowerPoint COM export unavailable" + (f": {detail[:240]}" if detail else "")


def export_with_libreoffice(pptx_path: Path, pdf_path: Path) -> tuple[bool, str]:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False, "LibreOffice export unavailable: soffice not found"
    try:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(pptx_path)],
            capture_output=True,
            text=True,
            timeout=120,
        )
    except Exception as exc:
        return False, f"LibreOffice export failed to start: {exc}"
    if result.returncode == 0 and pdf_path.exists():
        return True, "exported PDF via LibreOffice"
    detail = (result.stderr or result.stdout or "").strip().replace("\n", " ")
    return False, "LibreOffice export failed" + (f": {detail[:240]}" if detail else "")


def write_index(outputs: dict[str, Path], args: argparse.Namespace, source: str, spec: dict[str, Any], written: list[str], export_notes: list[str]) -> None:
    caption = spec.get("caption") or "图：PPT 思路图展示模型推理链条与执行路径。"
    interpretation = spec.get("interpretation") or "该图用于解释问题分解、模型模块和决策流向；每个节点与箭头应在正文中对应到变量、公式、算法步骤或执行动作。"
    lines = [
        "# PPT Reasoning Diagram Index",
        "",
        f"- Source: `{source}`",
        f"- Editable PPTX: `{outputs['pptx']}`",
        f"- Normalized spec: `{outputs['spec']}`",
        f"- Node table: `{outputs['nodes']}`",
        f"- Edge table: `{outputs['edges']}`",
        f"- Parameters: `{outputs['params']}`",
        f"- Node count: {len(spec.get('nodes', []))}",
        f"- Edge count: {len(spec.get('edges', []))}",
        "",
        "## Diagram Assets",
        "",
    ]
    for path in written:
        lines.append(f"- `{path}`")
    lines.extend(
        [
            "",
            "## Caption Draft",
            "",
            str(caption),
            "",
            "## Nearby Interpretation Draft",
            "",
            str(interpretation),
            "",
            "## Export Notes",
            "",
        ]
    )
    lines.extend([f"- {note}" for note in export_notes] or ["- PNG/PDF export completed or not requested."])
    lines.extend(
        [
            "",
            "## Caveat",
            "",
            "- PPTX is the editable source; use exported PNG/PDF in the final paper.",
            "- Do not use PPT diagrams as decoration; each node and arrow must correspond to a real model object, formula, result, or execution step.",
            "- If export is unavailable, open the PPTX in PowerPoint and export the slide before final-paper insertion.",
            "",
        ]
    )
    outputs["index"].write_text("\n".join(lines), encoding="utf-8")


def demo_spec(name: str) -> dict[str, Any]:
    if name == "pseudo_3d_architecture":
        return {
            "title": "伪 3D 分层建模与验证框架",
            "caption": "图：伪 3D 分层建模与验证框架用层叠厚度表示数据、模型、求解、验证和决策之间的语义层次。",
            "interpretation": "该图用于增强论文中的结构阅读：厚度只表示层、通道、模块容量或场景堆叠等真实含义；每个立体块必须能对应到变量表、模块表、算法步骤、结果表或验证证据。",
            "note": "PPTX 为可编辑源文件；最终论文插入导出的 PNG/PDF，并在正文说明深度维度的含义。",
            "groups": [
                {"label": "输入层", "x": 0.02, "y": 0.18, "w": 0.16, "h": 0.48, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#CBD5E1", "dashed": True, "text_color": "#6B7280"},
                {"label": "模型与求解层", "x": 0.22, "y": 0.12, "w": 0.46, "h": 0.60, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#9CCBC5", "dashed": True, "text_color": "#477C93"},
                {"label": "验证与决策层", "x": 0.72, "y": 0.18, "w": 0.25, "h": 0.50, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#D7B7DD", "dashed": True, "text_color": "#765187"},
            ],
            "nodes": [
                {"id": "data", "label": "数据张量/样本栈\nX", "type": "data", "x": 0.09, "y": 0.40, "w": 0.13, "h": 0.12, "fill": "#E8F2FB", "line": "#2F6B9A", "shape": "rect", "bold": True, "pseudo_3d": True, "depth_layers": 4, "depth_dx": 0.010, "depth_dy": -0.018},
                {"id": "features", "label": "特征层\nz=f(X)", "type": "model", "x": 0.27, "y": 0.30, "w": 0.14, "h": 0.11, "fill": "#DDEFEA", "line": "#2A9D8F", "bold": True, "pseudo_3d": True, "depth_layers": 3, "depth_dx": 0.012, "depth_dy": -0.016},
                {"id": "mechanism", "label": "机理/约束块\nF(x) <= b", "type": "equation", "x": 0.45, "y": 0.30, "w": 0.16, "h": 0.11, "fill": "#FFFFFF", "line": "#6B7280", "bold": True, "pseudo_3d": True, "depth_layers": 2, "depth_dx": 0.012, "depth_dy": -0.016},
                {"id": "solver", "label": "求解器模块\nDP / SA / MILP", "type": "model", "x": 0.61, "y": 0.45, "w": 0.16, "h": 0.12, "fill": "#EAF4EF", "line": "#2A9D8F", "bold": True, "pseudo_3d": True, "depth_layers": 3, "depth_dx": 0.011, "depth_dy": -0.017},
                {"id": "scenarios", "label": "场景方案栈\nS1...Sk", "type": "process", "x": 0.35, "y": 0.58, "w": 0.15, "h": 0.10, "fill": "#F4EFE8", "line": "#E76F51", "bold": True, "pseudo_3d": True, "depth_layers": 5, "depth_dx": 0.010, "depth_dy": -0.013},
                {"id": "validation", "label": "验证证据\n误差/敏感性", "type": "decision", "x": 0.78, "y": 0.34, "w": 0.14, "h": 0.12, "fill": "#FFF6D8", "line": "#C99A2E", "bold": True, "pseudo_3d": True, "depth_layers": 2, "depth_dx": 0.011, "depth_dy": -0.016},
                {"id": "decision", "label": "输出决策\n策略/建议", "type": "result", "x": 0.91, "y": 0.52, "w": 0.13, "h": 0.11, "fill": "#E9F5E6", "line": "#5E8C61", "bold": True, "pseudo_3d": True, "depth_layers": 2, "depth_dx": 0.010, "depth_dy": -0.016},
            ],
            "edges": [
                {"source": "data", "target": "features", "label": "清洗/编码"},
                {"source": "features", "target": "mechanism", "label": "变量"},
                {"source": "mechanism", "target": "solver", "label": "目标/约束"},
                {"source": "features", "target": "scenarios", "label": "场景", "connector": "elbow", "style": "secondary"},
                {"source": "scenarios", "target": "solver", "label": "候选"},
                {"source": "solver", "target": "validation", "label": "结果"},
                {"source": "validation", "target": "decision", "label": "通过"},
                {"source": "validation", "target": "mechanism", "label": "修正", "style": "feedback", "dashed": True, "lift": 0.20},
            ],
            "callouts": [
                {"target": "data", "text": "厚度=样本/通道/场景数量\n必须在图注或正文说明", "x": 0.05, "y": 0.69, "w": 0.22, "h": 0.10},
                {"target": "validation", "text": "伪 3D 只帮助阅读结构；\n数值结论仍需表格、公式或验证图。", "x": 0.67, "y": 0.72, "w": 0.28, "h": 0.11},
            ],
        }
    if name == "cnn_architecture":
        return {
            "title": "CNN 卷积特征提取与预测框架",
            "caption": "图：CNN 卷积特征提取与预测框架展示输入张量经过卷积、池化、向量化和输出头得到分类或回归结果的过程。",
            "interpretation": "该图用于说明卷积网络的结构证据：每个卷积块必须能对应到层参数表中的核大小、通道数、步长、填充和输出尺寸；预测头之后必须由训练设置、验证指标、基线对照或消融实验支撑。",
            "note": "PPTX 为可编辑源文件；最终论文插入导出的 PNG/PDF，并在正文附近放置层参数表与验证结果表。",
            "groups": [
                {"label": "输入与预处理", "x": 0.01, "y": 0.22, "w": 0.16, "h": 0.42, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#CBD5E1", "dashed": True, "text_color": "#6B7280"},
                {"label": "卷积特征提取", "x": 0.19, "y": 0.14, "w": 0.48, "h": 0.58, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#9CCBC5", "dashed": True, "text_color": "#477C93"},
                {"label": "预测头与证据", "x": 0.70, "y": 0.20, "w": 0.28, "h": 0.50, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#D7B7DD", "dashed": True, "text_color": "#765187"},
            ],
            "nodes": [
                {"id": "input", "label": "输入张量\n64x64x3", "type": "data", "x": 0.09, "y": 0.44, "w": 0.13, "h": 0.12, "fill": "#F7FAFC", "line": "#9AA8B5", "shape": "rect", "bold": True},
                {"id": "conv1", "label": "Conv 3x3\n32 通道", "type": "model", "x": 0.25, "y": 0.34, "w": 0.13, "h": 0.11, "fill": "#DDEFEA", "line": "#2A9D8F", "bold": True},
                {"id": "act1", "label": "BN + ReLU", "type": "process", "x": 0.25, "y": 0.50, "w": 0.12, "h": 0.075, "fill": "#F5EFE4", "line": "#C9845E", "font_size": 9},
                {"id": "pool1", "label": "MaxPool 2x2\n32x32", "type": "model", "x": 0.39, "y": 0.42, "w": 0.13, "h": 0.12, "fill": "#E8F2FB", "line": "#2F6B9A", "bold": True},
                {"id": "conv2", "label": "Conv 3x3\n64 通道", "type": "model", "x": 0.53, "y": 0.34, "w": 0.13, "h": 0.11, "fill": "#DDEFEA", "line": "#2A9D8F", "bold": True},
                {"id": "dropout", "label": "Dropout\np=0.2", "type": "process", "x": 0.53, "y": 0.50, "w": 0.12, "h": 0.075, "fill": "#F5EFE4", "line": "#C9845E", "font_size": 9},
                {"id": "pool2", "label": "Avg/MaxPool\n16x16", "type": "model", "x": 0.65, "y": 0.42, "w": 0.13, "h": 0.12, "fill": "#E8F2FB", "line": "#2F6B9A", "bold": True},
                {"id": "features", "label": "GAP/Flatten\n特征向量", "type": "equation", "x": 0.76, "y": 0.42, "w": 0.12, "h": 0.12, "fill": "#F8E8B8", "line": "#9B7A35", "bold": True},
                {"id": "dense", "label": "Dense + ReLU\n128 维", "type": "model", "x": 0.87, "y": 0.34, "w": 0.13, "h": 0.11, "fill": "#EAD5EC", "line": "#765187", "bold": True},
                {"id": "output", "label": "输出头\n分类/回归", "type": "result", "x": 0.95, "y": 0.50, "w": 0.10, "h": 0.10, "fill": "#F3E8F6", "line": "#765187", "bold": True},
                {"id": "metrics", "label": "验证证据\nbaseline / ablation", "type": "warning", "x": 0.84, "y": 0.66, "w": 0.18, "h": 0.08, "fill": "#FFF8E7", "line": "#D09225", "font_size": 8, "bold": True},
            ],
            "edges": [
                {"source": "input", "target": "conv1", "label": ""},
                {"source": "conv1", "target": "act1", "label": "", "connector": "elbow", "color": "#2A9D8F"},
                {"source": "act1", "target": "pool1", "label": "", "connector": "elbow", "color": "#2A9D8F"},
                {"source": "pool1", "target": "conv2", "label": ""},
                {"source": "conv2", "target": "dropout", "label": "", "connector": "elbow", "color": "#2A9D8F"},
                {"source": "dropout", "target": "pool2", "label": "", "connector": "elbow", "color": "#2A9D8F"},
                {"source": "pool2", "target": "features", "label": ""},
                {"source": "features", "target": "dense", "label": ""},
                {"source": "dense", "target": "output", "label": ""},
                {"source": "output", "target": "metrics", "label": "检验", "connector": "elbow", "color": "#D09225"},
            ],
            "callouts": [
                {"target": "conv2", "text": "正文字段：kernel/channel/stride/padding\n必须与层参数表一致", "x": 0.40, "y": 0.70, "w": 0.24, "h": 0.10},
                {"target": "metrics", "text": "contest_final 不允许只放结构图；\n需报告训练设置、指标和消融。", "x": 0.69, "y": 0.78, "w": 0.25, "h": 0.105},
            ],
        }
    if name == "multimodal_fusion":
        return {
            "title": "多源特征融合预测框架",
            "caption": "图：多源特征融合预测框架展示原始对象、序列特征、结构特征和统计属性经分支编码后进入融合模块并输出预测结果的过程。",
            "interpretation": "该图用于说明多分支模型结构：不同颜色代表不同特征来源，窄条表示分支编码后的潜在向量，虚线表示重构或辅助训练路径，融合模块之后的输出需要由验证指标和消融实验支撑。",
            "note": "PPTX 为可编辑源文件；最终论文插入导出的 PNG/PDF。",
            "groups": [
                {"label": "输入对象", "x": 0.01, "y": 0.18, "w": 0.16, "h": 0.50, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#CBD5E1", "dashed": True, "text_color": "#6B7280"},
                {"label": "分支特征编码", "x": 0.20, "y": 0.10, "w": 0.49, "h": 0.70, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#9CCBC5", "dashed": True, "text_color": "#477C93"},
                {"label": "融合与预测", "x": 0.71, "y": 0.17, "w": 0.27, "h": 0.48, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#D7B7DD", "dashed": True, "text_color": "#765187"},
            ],
            "nodes": [
                {"id": "object", "label": "原始对象\n/样本", "type": "data", "x": 0.08, "y": 0.42, "w": 0.12, "h": 0.13, "fill": "#F7FAFC", "line": "#9AA8B5", "shape": "ellipse", "bold": True},
                {"id": "seq_label", "label": "序列特征", "type": "title", "x": 0.23, "y": 0.18, "w": 0.12, "h": 0.07, "fill": "#DCECCB", "line": "#6E8C5B", "bold": True},
                {"id": "struct_label", "label": "结构特征", "type": "title", "x": 0.23, "y": 0.43, "w": 0.12, "h": 0.07, "fill": "#CBE3F6", "line": "#5B87A6", "bold": True},
                {"id": "prop_label", "label": "统计属性", "type": "title", "x": 0.23, "y": 0.68, "w": 0.12, "h": 0.07, "fill": "#F8E8B8", "line": "#9B7A35", "bold": True},
                {"id": "seq_features", "label": "x_seq\n编码序列", "type": "equation", "x": 0.36, "y": 0.18, "w": 0.13, "h": 0.055, "fill": "#DCECCB", "line": "#B7CFA2", "bold": True},
                {"id": "struct_features", "label": "G=(V,E)\n图结构", "type": "equation", "x": 0.36, "y": 0.43, "w": 0.13, "h": 0.075, "fill": "#CBE3F6", "line": "#89B6D8", "bold": True},
                {"id": "prop_features", "label": "z_prop\n指标向量", "type": "equation", "x": 0.36, "y": 0.68, "w": 0.13, "h": 0.075, "fill": "#F8E8B8", "line": "#D8BF73", "bold": True},
                {"id": "seq_encoder", "label": "F_seq\n序列编码器", "type": "model", "x": 0.51, "y": 0.18, "w": 0.13, "h": 0.10, "fill": "#EAF4D8", "line": "#5E8C61", "shape": "parallelogram", "bold": True},
                {"id": "struct_encoder", "label": "F_struct\n图网络", "type": "model", "x": 0.51, "y": 0.43, "w": 0.15, "h": 0.10, "fill": "#D9EDF9", "line": "#2F6B9A", "bold": True},
                {"id": "prop_encoder", "label": "F_prop\nMLP", "type": "model", "x": 0.51, "y": 0.68, "w": 0.13, "h": 0.08, "fill": "#FFF1C6", "line": "#9B7A35", "bold": True},
                {"id": "seq_latent", "label": "h_s", "type": "equation", "x": 0.63, "y": 0.18, "w": 0.035, "h": 0.13, "fill": "#DCECCB", "line": "#6E8C5B", "bold": True},
                {"id": "struct_latent", "label": "h_g", "type": "equation", "x": 0.63, "y": 0.43, "w": 0.035, "h": 0.13, "fill": "#CBE3F6", "line": "#2F6B9A", "bold": True},
                {"id": "prop_latent", "label": "h_p", "type": "equation", "x": 0.63, "y": 0.68, "w": 0.035, "h": 0.13, "fill": "#F8E8B8", "line": "#9B7A35", "bold": True},
                {"id": "decoder", "label": "G_seq\n重构器", "type": "model", "x": 0.60, "y": 0.07, "w": 0.12, "h": 0.08, "fill": "#EAF4D8", "line": "#5E8C61", "shape": "parallelogram", "bold": True},
                {"id": "merge_bar", "label": "", "type": "equation", "x": 0.72, "y": 0.43, "w": 0.035, "h": 0.48, "fill": "#E9DDF0", "line": "#765187"},
                {"id": "fusion", "label": "多模态\n注意力融合", "type": "model", "x": 0.82, "y": 0.43, "w": 0.13, "h": 0.23, "fill": "#EAD5EC", "line": "#765187", "bold": True},
                {"id": "head", "label": "预测头\nMLP/分类器", "type": "model", "x": 0.93, "y": 0.43, "w": 0.09, "h": 0.14, "fill": "#F5EAF6", "line": "#765187", "bold": True},
                {"id": "prediction", "label": "预测/评价\n结果", "type": "result", "x": 0.975, "y": 0.43, "w": 0.08, "h": 0.075, "fill": "#F3E8F6", "line": "#765187", "bold": True},
            ],
            "edges": [
                {"source": "object", "target": "seq_label", "label": "", "color": "#6E8C5B"},
                {"source": "object", "target": "struct_label", "label": "", "color": "#2F6B9A"},
                {"source": "object", "target": "prop_label", "label": "", "color": "#9B7A35"},
                {"source": "seq_label", "target": "seq_features", "label": ""},
                {"source": "struct_label", "target": "struct_features", "label": ""},
                {"source": "prop_label", "target": "prop_features", "label": ""},
                {"source": "seq_features", "target": "seq_encoder", "label": ""},
                {"source": "struct_features", "target": "struct_encoder", "label": ""},
                {"source": "prop_features", "target": "prop_encoder", "label": ""},
                {"source": "seq_encoder", "target": "seq_latent", "label": ""},
                {"source": "struct_encoder", "target": "struct_latent", "label": ""},
                {"source": "prop_encoder", "target": "prop_latent", "label": ""},
                {"source": "seq_latent", "target": "decoder", "label": "辅助重构", "style": "feedback", "color": "#9B2C2C", "dashed": True, "lift": 0.13},
                {"source": "decoder", "target": "seq_features", "label": "reconstruction", "style": "feedback", "color": "#9B2C2C", "dashed": True, "lift": 0.12},
                {"source": "seq_latent", "target": "merge_bar", "label": "", "connector": "elbow", "color": "#6E8C5B"},
                {"source": "struct_latent", "target": "merge_bar", "label": "", "connector": "elbow", "color": "#2F6B9A"},
                {"source": "prop_latent", "target": "merge_bar", "label": "", "connector": "elbow", "color": "#9B7A35"},
                {"source": "merge_bar", "target": "fusion", "label": ""},
                {"source": "fusion", "target": "head", "label": ""},
                {"source": "head", "target": "prediction", "label": ""},
            ],
            "callouts": [
                {"target": "fusion", "text": "融合层需要\n权重/注意力公式\n或消融实验支撑", "x": 0.74, "y": 0.70, "w": 0.20, "h": 0.10},
            ],
        }
    if name == "showcase_work":
        return {
            "title": "模型工作框架总览",
            "caption": "图：模型工作框架总览展示数据输入、核心模型、敏感性分析、优化决策和最终目标之间的逻辑关系。",
            "interpretation": "该图采用优秀论文中常见的总览式思路图语法：左侧给出输入与预处理，中央展示主模型链条，右侧列出求解条件与控制模块，中部展开敏感性分支，底部收束到优化目标。",
            "note": "虚线框表示同一建模阶段；箭头表示数据、参数或策略的传递关系。",
            "groups": [
                {"label": "", "x": 0.17, "y": 0.04, "w": 0.80, "h": 0.90, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#7CB7B0", "line_width": 1.0, "dashed": True},
                {"label": "核心机理模型", "x": 0.185, "y": 0.105, "w": 0.36, "h": 0.40, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#9CCBC5", "dashed": True, "text_color": "#477C93"},
                {"label": "动态求解模块", "x": 0.575, "y": 0.105, "w": 0.37, "h": 0.40, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#CBD5E1", "dashed": True, "text_color": "#6B7280"},
                {"label": "敏感性分析", "x": 0.185, "y": 0.535, "w": 0.76, "h": 0.20, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#CBD5E1", "dashed": True, "text_color": "#477C93"},
                {"label": "目标与优化", "x": 0.185, "y": 0.765, "w": 0.76, "h": 0.17, "shape": "rect", "fill": "#FFFFFF", "transparency": 100, "line": "#D8A798", "dashed": True, "text_color": "#B56B56"},
            ],
            "nodes": [
                {"id": "data", "label": "原始数据", "type": "data", "x": 0.07, "y": 0.70, "w": 0.12, "h": 0.065, "fill": "#B9CDEF", "line": "#5B7EAF", "bold": True},
                {"id": "processing", "label": "数据处理", "type": "data", "x": 0.07, "y": 0.54, "w": 0.12, "h": 0.065, "fill": "#D8E8F4", "line": "#8CA9C4", "bold": True},
                {"id": "goal", "label": "基于多源信息的综合决策建模", "type": "title", "x": 0.56, "y": 0.085, "w": 0.72, "h": 0.055, "bold": True, "font_size": 9},
                {"id": "base_model", "label": "三参数基础模型", "type": "equation", "x": 0.36, "y": 0.19, "w": 0.31, "h": 0.055, "bold": True},
                {"id": "mechanism", "label": "机理修正模型", "type": "model", "x": 0.36, "y": 0.29, "w": 0.31, "h": 0.055},
                {"id": "solution", "label": "动态方程求解", "type": "model", "x": 0.36, "y": 0.37, "w": 0.31, "h": 0.055},
                {"id": "profile", "label": "特征剖面与分布", "type": "title", "x": 0.36, "y": 0.49, "w": 0.33, "h": 0.055, "bold": True, "font_size": 9},
                {"id": "dynamic", "label": "动态解", "type": "equation", "x": 0.76, "y": 0.18, "w": 0.34, "h": 0.052},
                {"id": "control", "label": "最优控制", "type": "equation", "x": 0.76, "y": 0.25, "w": 0.34, "h": 0.052},
                {"id": "condition", "label": "必要条件", "type": "equation", "x": 0.76, "y": 0.32, "w": 0.34, "h": 0.052},
                {"id": "dp", "label": "动态规划", "type": "equation", "x": 0.76, "y": 0.39, "w": 0.34, "h": 0.052},
                {"id": "principle", "label": "递推原则", "type": "equation", "x": 0.76, "y": 0.46, "w": 0.34, "h": 0.052},
                {"id": "weather", "label": "环境\n敏感性", "type": "title", "x": 0.28, "y": 0.63, "w": 0.13, "h": 0.10, "font_size": 9, "bold": True},
                {"id": "wind", "label": "风速影响", "type": "equation", "x": 0.48, "y": 0.59, "w": 0.22, "h": 0.055},
                {"id": "temp", "label": "温度影响", "type": "equation", "x": 0.48, "y": 0.68, "w": 0.22, "h": 0.055},
                {"id": "deviation", "label": "偏差\n敏感性", "type": "title", "x": 0.68, "y": 0.63, "w": 0.13, "h": 0.10, "font_size": 9, "bold": True},
                {"id": "prob", "label": "偏差概率", "type": "equation", "x": 0.86, "y": 0.59, "w": 0.20, "h": 0.055},
                {"id": "loc", "label": "位置偏差", "type": "equation", "x": 0.86, "y": 0.68, "w": 0.20, "h": 0.055},
                {"id": "team_goal", "label": "综合优化目标", "type": "title", "x": 0.43, "y": 0.80, "w": 0.38, "h": 0.055, "bold": True},
                {"id": "opt_model", "label": "优化模型", "type": "equation", "x": 0.78, "y": 0.80, "w": 0.28, "h": 0.055},
                {"id": "reduce", "label": "降低代价", "type": "equation", "x": 0.31, "y": 0.90, "w": 0.19, "h": 0.055, "bold": True},
                {"id": "maximize", "label": "提升系统收益", "type": "equation", "x": 0.55, "y": 0.90, "w": 0.23, "h": 0.055, "bold": True},
            ],
            "edges": [
                {"source": "data", "target": "processing", "label": ""},
                {"source": "processing", "target": "base_model", "label": "", "connector": "elbow"},
                {"source": "goal", "target": "base_model", "label": "", "connector": "elbow"},
                {"source": "base_model", "target": "mechanism", "label": ""},
                {"source": "mechanism", "target": "solution", "label": ""},
                {"source": "solution", "target": "profile", "label": ""},
                {"source": "solution", "target": "dynamic", "label": "", "connector": "elbow"},
                {"source": "dynamic", "target": "control", "label": ""},
                {"source": "control", "target": "condition", "label": ""},
                {"source": "condition", "target": "dp", "label": ""},
                {"source": "dp", "target": "principle", "label": ""},
                {"source": "profile", "target": "weather", "style": "secondary", "label": "", "connector": "elbow"},
                {"source": "weather", "target": "wind", "style": "secondary", "label": "", "connector": "elbow"},
                {"source": "weather", "target": "temp", "style": "secondary", "label": "", "connector": "elbow"},
                {"source": "profile", "target": "deviation", "style": "secondary", "label": "", "connector": "elbow"},
                {"source": "deviation", "target": "prob", "style": "secondary", "label": "", "connector": "elbow"},
                {"source": "deviation", "target": "loc", "style": "secondary", "label": "", "connector": "elbow"},
                {"source": "profile", "target": "team_goal", "label": "", "connector": "elbow"},
                {"source": "team_goal", "target": "opt_model", "label": ""},
                {"source": "team_goal", "target": "reduce", "label": ""},
                {"source": "team_goal", "target": "maximize", "label": ""},
            ],
        }
    if name == "decision_loop":
        return {
            "title": "稳健决策闭环思路图",
            "caption": "图：从数据更新到稳健决策执行的闭环流程。",
            "interpretation": "该图把估计、优化、阈值判断、执行和复盘连成闭环，用于说明模型结果如何转化为可执行策略。",
            "note": "虚线表示反馈更新；实线表示主决策流。",
            "groups": [
                {"label": "数据层", "x": 0.02, "y": 0.10, "w": 0.28, "h": 0.78, "type": "data"},
                {"label": "模型层", "x": 0.34, "y": 0.10, "w": 0.32, "h": 0.78, "type": "model"},
                {"label": "执行层", "x": 0.70, "y": 0.10, "w": 0.28, "h": 0.78, "type": "result"},
            ],
            "nodes": [
                {"id": "data", "label": "样本数据\n异常清洗", "type": "data"},
                {"id": "estimate", "label": "参数估计\n置信区间", "type": "model"},
                {"id": "opt", "label": "优化求解\n候选策略", "type": "model"},
                {"id": "threshold", "label": "阈值判断", "type": "decision"},
                {"id": "execute", "label": "执行方案\n责任人", "type": "result"},
                {"id": "review", "label": "KPI复盘\n数据回流", "type": "process"},
            ],
            "edges": [
                {"source": "data", "target": "estimate", "label": "输入"},
                {"source": "estimate", "target": "opt", "label": "参数"},
                {"source": "opt", "target": "threshold", "label": "收益/风险"},
                {"source": "threshold", "target": "execute", "label": "通过"},
                {"source": "execute", "target": "review", "label": "运行"},
                {"source": "review", "target": "data", "label": "更新", "style": "feedback", "dashed": True},
            ],
            "callouts": [
                {"target": "threshold", "text": "阈值条件必须由\n灵敏度或约束审计支撑"},
            ],
        }
    if name == "paper_logic":
        return {
            "title": "多问题论文推理链条",
            "caption": "图：四个子问题之间的数据、模型和结论传递关系。",
            "interpretation": "该图用于约束论文结构，避免各问割裂；后一问的模型输入必须能追溯到前一问的结果或假设。",
            "nodes": [
                {"id": "q1", "label": "问题一\n数据识别", "type": "data"},
                {"id": "q2", "label": "问题二\n基础模型", "type": "model"},
                {"id": "q3", "label": "问题三\n扩展优化", "type": "model"},
                {"id": "q4", "label": "问题四\n稳健检验", "type": "decision"},
                {"id": "final", "label": "最终建议\n管理方案", "type": "result"},
            ],
            "edges": [
                {"source": "q1", "target": "q2", "label": "参数"},
                {"source": "q2", "target": "q3", "label": "状态/约束"},
                {"source": "q3", "target": "q4", "label": "候选方案"},
                {"source": "q4", "target": "final", "label": "可信策略"},
            ],
            "callouts": [
                {"target": "q3", "text": "扩展问要保留\n前问符号和约束"},
                {"target": "final", "text": "结论必须连接\n数值结果与执行动作"},
            ],
        }
    return {
        "title": "建模求解总思路图",
        "caption": "图：从题目解析到最终建议的建模求解流程。",
        "interpretation": "该图把数据、模型、求解、验证和论文结论的对应关系显式化，便于正文按图中的链条展开。",
        "groups": [
            {"label": "输入与假设", "x": 0.02, "y": 0.13, "w": 0.25, "h": 0.70, "type": "data"},
            {"label": "建模与求解", "x": 0.30, "y": 0.13, "w": 0.42, "h": 0.70, "type": "model"},
            {"label": "验证与输出", "x": 0.75, "y": 0.13, "w": 0.23, "h": 0.70, "type": "result"},
        ],
        "nodes": [
            {"id": "problem", "label": "题目解析\n指标提取", "type": "data"},
            {"id": "assumption", "label": "假设与变量\n符号表", "type": "equation"},
            {"id": "model", "label": "机制模型\n目标函数", "type": "model"},
            {"id": "solver", "label": "算法求解\n复杂度说明", "type": "process"},
            {"id": "validation", "label": "误差/稳健性\n对照检验", "type": "decision"},
            {"id": "answer", "label": "结果冻结\n论文结论", "type": "result"},
        ],
        "edges": [
            {"source": "problem", "target": "assumption", "label": "约束"},
            {"source": "assumption", "target": "model", "label": "变量"},
            {"source": "model", "target": "solver", "label": "目标/约束"},
            {"source": "solver", "target": "validation", "label": "候选解"},
            {"source": "validation", "target": "answer", "label": "可信结果"},
            {"source": "validation", "target": "model", "label": "修正", "style": "feedback", "dashed": True},
        ],
        "callouts": [
            {"target": "model", "text": "核心公式附近\n应放置对应图示"},
            {"target": "validation", "text": "验证失败则返回\n模型或算法阶段"},
        ],
    }


def style_for(kind: Any) -> dict[str, str]:
    return NODE_STYLES.get(str(kind).lower(), NODE_STYLES["default"])


def edge_color(edge: dict[str, Any]) -> str:
    style = str(edge.get("style", "main")).lower()
    if style == "feedback":
        return COLORS["red"]
    if style in {"secondary", "optional"}:
        return COLORS["teal"]
    if style == "assumption":
        return COLORS["gray"]
    return COLORS["blue"]


def to_box(x: float, y: float, w: float, h: float) -> tuple[Any, Any, Any, Any]:
    cw = SLIDE_W - CANVAS["left"] - CANVAS["right"]
    ch = SLIDE_H - CANVAS["top"] - CANVAS["bottom"]
    return (
        Inches(CANVAS["left"] + x * cw),
        Inches(CANVAS["top"] + y * ch),
        Inches(w * cw),
        Inches(h * ch),
    )


def to_point(x: float, y: float) -> tuple[Any, Any]:
    cw = SLIDE_W - CANVAS["left"] - CANVAS["right"]
    ch = SLIDE_H - CANVAS["top"] - CANVAS["bottom"]
    return Inches(CANVAS["left"] + x * cw), Inches(CANVAS["top"] + y * ch)


def rgb(value: str) -> RGBColor:
    value = value.strip().lstrip("#")
    if len(value) != 6:
        value = "1F2933"
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames: list[str] = []
    for row in rows:
        for key in row.keys():
            if key not in fieldnames:
                fieldnames.append(key)
    if not fieldnames:
        fieldnames = ["empty"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def prepare_outputs(args: argparse.Namespace, root: Path) -> dict[str, Path]:
    diagrams_dir = resolve_path(root, args.diagrams_dir)
    data_dir = resolve_path(root, args.data_dir)
    diagrams_dir.mkdir(parents=True, exist_ok=True)
    data_dir.mkdir(parents=True, exist_ok=True)
    prefix = safe_prefix(args.prefix)
    stem = f"{prefix}_ppt_reasoning"
    return {
        "diagrams_dir": diagrams_dir,
        "data_dir": data_dir,
        "pptx": diagrams_dir / f"{stem}.pptx",
        "png": diagrams_dir / f"{stem}.png",
        "pdf": diagrams_dir / f"{stem}.pdf",
        "spec": data_dir / f"{stem}_spec.json",
        "nodes": data_dir / f"{stem}_nodes.csv",
        "edges": data_dir / f"{stem}_edges.csv",
        "groups": data_dir / f"{stem}_groups.csv",
        "callouts": data_dir / f"{stem}_callouts.csv",
        "params": data_dir / f"{stem}_params.json",
        "index": diagrams_dir / f"{stem}_index.md",
    }


def safe_prefix(value: str) -> str:
    prefix = re.sub(r"[^0-9A-Za-z_\-\u4e00-\u9fff]+", "_", value.strip())
    return prefix or "ppt_reasoning"


def ps_quote(path: Path) -> str:
    return "'" + str(path).replace("'", "''") + "'"


def resolve_path(root: Path, value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else root / path


def clamp(value: float, low: float, high: float) -> float:
    return max(low, min(high, value))


def clamp_int(value: int, low: int, high: int) -> int:
    return max(low, min(high, value))


if __name__ == "__main__":
    raise SystemExit(main())
